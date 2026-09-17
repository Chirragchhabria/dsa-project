#!/usr/bin/env python3
"""Generate the verified midterm report from the repository's implemented state."""

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "SOC_Alert_Correlation_Midterm_Report.docx"
ASSET_DIR = ROOT / "build" / "report_assets"
DIAGRAM = ASSET_DIR / "architecture.png"

NAVY = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GRAY = "F5F7F9"
MID_GRAY = "D9D9D9"
TEXT_GRAY = RGBColor(70, 70, 70)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color=MID_GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_heading(paragraph, keep_next=True):
    p_pr = paragraph._p.get_or_add_pPr()
    keep = OxmlElement("w:keepNext") if keep_next else OxmlElement("w:keepLines")
    p_pr.append(keep)


def set_run_font(run, name="Arial"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)


def add_toc_line(doc, title, page, indent=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(indent)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(6.65 - indent), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = paragraph.add_run(title + "\t" + str(page))
    set_run_font(run)
    run.font.size = Pt(10.2)


def add_paragraph(doc, text="", bold_lead=None, style=None):
    paragraph = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        lead.bold = True
        set_run_font(lead)
        body = paragraph.add_run(text[len(bold_lead):])
        set_run_font(body)
    else:
        run = paragraph.add_run(text)
        set_run_font(run)
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Number")
        run = paragraph.add_run(item)
        set_run_font(run)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    set_repeat_heading(paragraph)
    return paragraph


def add_table(doc, headers, rows, widths=None, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        if widths:
            cell.width = Inches(widths[index])
        set_cell_shading(cell, NAVY)
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(header)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(font_size)
        set_run_font(run)

    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        prevent_row_split(table.rows[-1])
        for index, value in enumerate(values):
            cell = cells[index]
            if widths:
                cell.width = Inches(widths[index])
            set_cell_borders(cell)
            set_cell_margins(cell)
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
            set_run_font(run)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def create_architecture_diagram():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    width, height = 1800, 980
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    regular = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 34)
    small = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial.ttf", 27)
    bold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Arial Bold.ttf", 36)

    def box(x1, y1, x2, y2, title, subtitle, fill):
        draw.rounded_rectangle((x1, y1, x2, y2), radius=20, fill=fill, outline="#17365D", width=4)
        bbox = draw.textbbox((0, 0), title, font=bold)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y1 + 24), title, font=bold, fill="#000000")
        lines = subtitle.split("\n")
        y = y1 + 84
        for line in lines:
            line_box = draw.textbbox((0, 0), line, font=small)
            draw.text(((x1 + x2 - (line_box[2] - line_box[0])) / 2, y), line, font=small, fill="#333333")
            y += 34

    def arrow(x1, y1, x2, y2):
        draw.line((x1, y1, x2, y2), fill="#17365D", width=6)
        if abs(x2 - x1) > abs(y2 - y1):
            direction = 1 if x2 > x1 else -1
            points = [(x2, y2), (x2 - 22 * direction, y2 - 14), (x2 - 22 * direction, y2 + 14)]
        else:
            direction = 1 if y2 > y1 else -1
            points = [(x2, y2), (x2 - 14, y2 - 22 * direction), (x2 + 14, y2 - 22 * direction)]
        draw.polygon(points, fill="#17365D")

    box(60, 80, 360, 250, "CSV Input", "validated security\nevent records", "#F5F7F9")
    box(490, 80, 870, 250, "EventManager", "event store and\nhash map indexes", "#EAF2F8")
    box(1000, 80, 1390, 250, "SlidingWindow", "recent events within\n300 seconds", "#EAF2F8")
    box(1000, 390, 1390, 570, "CorrelationEngine", "entity matching and\nevent family rules", "#DDEBF7")
    box(490, 390, 870, 570, "EntityGraph", "adjacency list with\nBFS and recursive DFS", "#DDEBF7")
    box(490, 710, 870, 890, "Incident", "linked event timeline\nand transparent score", "#E2F0D9")
    box(1000, 710, 1390, 890, "PriorityQueue", "highest risk incident\nreturned first", "#E2F0D9")
    box(1510, 710, 1740, 890, "CLI", "analyst\noutput", "#FCE4D6")

    arrow(360, 165, 490, 165)
    arrow(870, 165, 1000, 165)
    arrow(1195, 250, 1195, 390)
    arrow(1000, 480, 870, 480)
    arrow(680, 570, 680, 710)
    arrow(870, 800, 1000, 800)
    arrow(1390, 800, 1510, 800)

    draw.text((60, 910), "Non-owning event references are used during processing; incidents store owned event snapshots.",
              font=regular, fill="#333333")
    image.save(DIAGRAM, dpi=(200, 200))


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(12)
    title_p_pr = title._element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    for style_name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(3)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run("SOC Alert Correlation and Incident Prioritization Engine")
    set_run_font(run)
    run.font.size = Pt(8.5)
    run.font.color.rgb = TEXT_GRAY

    footer = section.footer
    add_page_number(footer.paragraphs[0])
    for run in footer.paragraphs[0].runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = TEXT_GRAY

    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("SOC Alert Correlation and Incident Prioritization Engine")
    set_run_font(run)
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    p_pr = title._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(8)
    subtitle.paragraph_format.space_after = Pt(30)
    run = subtitle.add_run("Midterm Project Report")
    set_run_font(run)
    run.font.size = Pt(16)
    run.font.bold = True

    course = doc.add_paragraph()
    course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = course.add_run("Data Structures Using C++")
    set_run_font(run)
    run.font.size = Pt(12)

    doc.add_paragraph()
    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.line_spacing = 1.5
    for index, line in enumerate(("Harish M", "Registration Number 230932300", "", "Chirrag Chhabria", "Registration Number 230932288")):
        run = authors.add_run(line)
        set_run_font(run)
        run.font.size = Pt(11.5)
        if line in ("Harish M", "Chirrag Chhabria"):
            run.bold = True
        if index < 4:
            run.add_break()

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(55)
    run = date.add_run("September 2026")
    set_run_font(run)
    run.font.size = Pt(11)

    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_before = Pt(34)
    run = status.add_run("Implementation status verified against repository commit df8ed43")
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = TEXT_GRAY
    doc.add_page_break()


def build_report():
    create_architecture_diagram()
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "SOC Alert Correlation and Incident Prioritization Engine Midterm Project Report"
    doc.core_properties.author = "Harish M and Chirrag Chhabria"
    doc.core_properties.subject = "Data Structures Using C++ midterm project"
    doc.core_properties.comments = "Generated from verified repository state"

    add_cover(doc)

    add_heading(doc, "Table of Contents", 1)
    for toc_title, page, indent in (
        ("Abstract", 3, 0),
        ("1 Problem Definition", 3, 0),
        ("2 Project Objectives and Scope", 3, 0),
        ("3 System Architecture", 3, 0),
        ("4 Module Design", 4, 0),
        ("5 Data Structures and Algorithms", 5, 0),
        ("6 Event Ingestion and Validation", 6, 0),
        ("7 Correlation Method", 6, 0),
        ("8 Incident Scoring and Prioritization", 6, 0),
        ("9 Synthetic Dataset", 7, 0),
        ("10 Testing and Verification", 7, 0),
        ("11 Demonstration Results", 8, 0),
        ("12 Complexity Analysis", 9, 0),
        ("13 Current Limitations", 9, 0),
        ("14 Planned Work", 10, 0),
        ("15 Conclusion", 10, 0),
        ("Appendix A Repository and Execution", 10, 0),
        ("Appendix B Verified Test Summary", 10, 0),
    ):
        add_toc_line(doc, toc_title, page, indent)
    doc.add_page_break()

    add_heading(doc, "Abstract", 1)
    add_paragraph(doc,
        "Security Operations Centers receive alerts from authentication, endpoint, network, and other security systems. "
        "A single alert often provides too little context, while repeated or related alerts increase analyst workload when reviewed independently. "
        "This project implements a modular C++ engine that groups temporally and logically related security events into incidents and ranks those incidents for review. "
        "The midterm version combines hash indexes, a custom linked list, a sliding-window queue, an adjacency-list graph, iterative breadth-first search, recursive depth-first search, and a priority queue. "
        "A controlled dataset of 23 events produced three incidents from three planted defensive scenarios. Seven behavioral test groups passed, including a sanitizer-enabled run. "
        "The current results establish functional correctness for the milestone; they do not establish production security accuracy or performance.")

    add_heading(doc, "1 Problem Definition", 1)
    add_paragraph(doc,
        "Security monitoring tools usually report individual observations: a failed login, a connection attempt, a privilege change, or a file access. "
        "Analysts must determine whether several observations describe one sequence and whether that sequence deserves immediate attention. "
        "Reviewing every event as an independent case hides context and creates avoidable queue volume.")
    add_paragraph(doc,
        "The implemented system addresses two questions. First, which events have sufficient shared context to belong to one incident? "
        "Second, which generated incident should an analyst examine first? The engine answers these questions with conservative, inspectable rules rather than an opaque probability model.")

    add_heading(doc, "2 Project Objectives and Scope", 1)
    add_heading(doc, "2.1 Midterm Objectives", 2)
    add_bullets(doc, [
        "Ingest security events from CSV and reject malformed records without discarding valid records.",
        "Index events by identifier, source IP, destination IP, user, and host for fast lookup.",
        "Maintain recent events in a configurable temporal window.",
        "Represent event and entity relationships in an adjacency-list graph.",
        "Use BFS and recursive DFS for graph exploration and connected-component discovery.",
        "Create chronological incident timelines with a manually implemented linked list.",
        "Assign explainable risk scores and order incidents with a max heap.",
        "Demonstrate the complete pipeline with controlled data and behavioral tests.",
    ])
    add_heading(doc, "2.2 Scope Boundaries", 2)
    add_paragraph(doc,
        "The midterm system is a defensive event-correlation prototype. It does not scan networks, exploit systems, or execute response actions. "
        "It also does not claim statistically validated detection accuracy, production throughput, or memory efficiency. "
        "Those claims require larger datasets, ground-truth evaluation, and repeatable benchmarks that have not yet been completed.")

    add_heading(doc, "3 System Architecture", 1)
    add_paragraph(doc,
        "The engine separates ingestion, temporal processing, relationship modeling, incident construction, and prioritization. "
        "This separation keeps each data structure tied to a specific operational responsibility and allows later rule or scoring improvements without replacing the storage layer.")
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    architecture_picture = picture.add_run().add_picture(str(DIAGRAM), width=Inches(6.72))
    architecture_picture._inline.docPr.set(
        "descr",
        "Processing flow from validated CSV input through EventManager, SlidingWindow, CorrelationEngine, EntityGraph, Incident, PriorityQueue, and CLI output",
    )
    architecture_picture._inline.docPr.set("title", "Midterm processing architecture")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run("Figure 1  Midterm processing architecture")
    set_run_font(run)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = TEXT_GRAY

    add_heading(doc, "3.1 Ownership and Data Flow", 2)
    add_paragraph(doc,
        "EventManager owns the ingested SecurityEvent values. Its indexes store vector positions instead of owning pointers. "
        "SlidingWindow temporarily holds non-owning pointers while the manager remains alive. EventLinkedList owns its nodes through unique pointers and stores event snapshots inside each incident. "
        "IncidentPrioritizer holds non-owning incident pointers while the incident collection remains in scope. These rules make lifetime and ownership explicit.")

    add_heading(doc, "4 Module Design", 1)
    add_table(doc,
        ["Module", "Current responsibility"],
        [
            ("SecurityEvent", "Stores the typed event record, severity, entities, description, source, and optional ground-truth label."),
            ("EventManager", "Loads and validates CSV rows, owns events, prevents duplicate identifiers, and maintains hash indexes."),
            ("EventLinkedList", "Stores an incident timeline in chronological order using manually allocated nodes owned by unique pointers."),
            ("SlidingWindow", "Keeps only events that fall within the configured temporal distance."),
            ("EntityGraph", "Stores event and entity nodes in an adjacency list and provides BFS, recursive DFS, and event components."),
            ("CorrelationEngine", "Creates time-bounded direct event edges when entity and event-family rules pass."),
            ("Incident", "Collects correlated events and entities, calculates the time range, and produces an explained score."),
            ("IncidentPrioritizer", "Returns the highest-risk incident first through a priority queue."),
        ],
        widths=[1.62, 5.0], font_size=9.2)

    add_heading(doc, "5 Data Structures and Algorithms", 1)
    add_table(doc,
        ["Concept", "Implementation", "Why it is used"],
        [
            ("Hash maps", "unordered_map indexes", "Provide average constant-time access to events by identifiers and security entities."),
            ("Queue", "deque in SlidingWindow", "Expires old events from the front while new events arrive at the back."),
            ("Linked list", "Custom Node and unique_ptr chain", "Maintains a chronological, owned incident timeline and demonstrates explicit node management."),
            ("Graph", "unordered_map of adjacency sets", "Represents event, IP, user, and host relationships without a dense matrix."),
            ("BFS", "Iterative queue traversal", "Explores reachable graph nodes without recursion."),
            ("DFS", "Recursive traversal", "Discovers connected security activity and event components with genuine recursion."),
            ("Priority queue", "std::priority_queue max heap", "Returns the highest risk score first with deterministic tie breaking."),
            ("Sorting", "stable_sort and sort", "Orders input by time, timelines chronologically, and component output deterministically."),
        ],
        widths=[1.15, 2.15, 3.32], font_size=8.8)

    add_heading(doc, "5.1 Graph Traversal", 2)
    add_paragraph(doc,
        "BFS uses an explicit queue and visited set. Neighbor keys are sorted before insertion into the queue so demonstrations remain deterministic. "
        "DFS calls a recursive helper for each unvisited neighbor. Event-component discovery uses a second recursive traversal that follows only direct event-to-event edges. "
        "Entity nodes remain available for exploration but are excluded from incident components because traversing through a common account or server could merge unrelated activity.")

    add_heading(doc, "6 Event Ingestion and Validation", 1)
    add_paragraph(doc,
        "The CSV reader supports quoted fields and escaped quotation marks. It requires event_id, timestamp, event_type, and severity columns. "
        "Timestamps must be non-negative Unix seconds, identifiers must be unique, and every event must contain at least one correlatable IP, user, or host. "
        "Unknown event types and severity values are rejected. Each rejected row produces a line-specific error while subsequent valid rows continue to load.")
    add_paragraph(doc,
        "The scenario_label column is optional. It records planted ground truth for later evaluation and is never consulted by the correlation or scoring algorithms.")

    add_heading(doc, "7 Correlation Method", 1)
    add_paragraph(doc,
        "Events are sorted by timestamp before processing. For each event, the engine examines only records still present in the five-minute sliding window. "
        "It creates a direct event edge when the records share an IP address, username, or hostname and belong to the same supported event family. Connected components with at least two events become incidents.")
    add_table(doc,
        ["Event family", "Included event types"],
        [
            ("Credential compromise", "Failed login, successful login, privilege escalation, and file access"),
            ("Reconnaissance", "Connection attempt, port scan, and service discovery"),
            ("Possible exfiltration", "Privileged access, sensitive file access, outbound connection, and large outbound transfer"),
        ],
        widths=[1.9, 4.72], font_size=9.3)
    add_paragraph(doc,
        "Normal activity and unknown event types do not form correlation edges in version 1. This conservative rule prevents ordinary events from joining an incident solely because they share infrastructure.")

    add_heading(doc, "8 Incident Scoring and Prioritization", 1)
    add_paragraph(doc,
        "The score is a deterministic prioritization heuristic. Each contribution is printed in the CLI so an analyst or evaluator can reproduce the result. "
        "The final value is capped at 100 and is not presented as a probability.")
    add_table(doc,
        ["Factor", "Contribution"],
        [
            ("Event severity", "Informational 1, Low 3, Medium 7, High 12, Critical 20 per event"),
            ("Chain length", "2 points per additional event, capped at 20"),
            ("Privilege activity", "15 points when privilege escalation or privileged access occurs"),
            ("Sensitive access or large transfer", "20 points when sensitive file access or a large outbound transfer occurs"),
            ("Affected hosts", "3 points per additional host, capped at 12"),
            ("Temporal density", "10 points for at least three events completed within five minutes"),
        ],
        widths=[2.1, 4.52], font_size=9.2)
    add_paragraph(doc,
        "Priority thresholds are Informational below 10, Low from 10 to 24, Medium from 25 to 49, High from 50 to 74, and Critical from 75 to 100. "
        "The max heap orders by risk score, then by earlier start time, then by incident identifier.")

    add_heading(doc, "9 Synthetic Dataset", 1)
    add_paragraph(doc,
        "The checked-in dataset contains 23 defensive simulation events: 10 normal records and 13 records across three planted scenarios. "
        "The scenarios exercise different relationship patterns while remaining safe and non-operational.")
    add_table(doc,
        ["Scenario", "Event sequence", "Records"],
        [
            ("Credential compromise", "Repeated failed login, successful login, privilege escalation, restricted file access", "5"),
            ("Reconnaissance", "Connection attempts, scan-like activity, service discovery", "4"),
            ("Possible exfiltration", "Privileged access, sensitive file access, unusual connection, large transfer", "4"),
            ("Normal activity", "Routine sign-ins, application access, synchronization, updates, and health checks", "10"),
        ],
        widths=[1.68, 4.15, 0.79], font_size=8.9)

    add_heading(doc, "10 Testing and Verification", 1)
    add_paragraph(doc,
        "The test executable checks behavior rather than trivial accessors. It was compiled with Apple Clang 21.0.0 using C++17, strict warnings, AddressSanitizer, and UndefinedBehaviorSanitizer. "
        "Both the normal and sanitizer-enabled executions reported seven passing test groups and zero failures.")
    add_table(doc,
        ["Test group", "Verified behavior", "Result"],
        [
            ("CSV parsing", "Quoted commas, malformed timestamps, duplicate identifiers, and row-level rejection", "Pass"),
            ("Hash indexes", "Source IP, user, host, and unique-entity counts", "Pass"),
            ("Linked list", "Chronological insertion, lookup, tail removal, deep copy, and clear", "Pass"),
            ("Sliding window", "Expiration beyond the configured duration", "Pass"),
            ("Graph traversal", "Edges, node and edge counts, BFS, recursive DFS, and event components", "Pass"),
            ("Correlation", "Expected chain creation plus rejection of unrelated and temporally distant events", "Pass"),
            ("Priority queue", "Higher scored incident returned before lower scored incident", "Pass"),
        ],
        widths=[1.38, 4.42, 0.82], font_size=8.7)
    add_paragraph(doc,
        "CMake target definitions are included in the repository. CMake was not installed in the verification environment, so the report does not claim that CMake or CTest was executed. "
        "Equivalent compile and link commands were run directly with Clang.")

    add_heading(doc, "11 Demonstration Results", 1)
    add_paragraph(doc,
        "The command-line demonstration loaded the checked-in dataset without rejected rows and generated three incidents. "
        "All values in this section came from the actual program execution on 17 September 2026.")
    add_table(doc,
        ["Metric", "Observed value"],
        [
            ("Events loaded", "23"),
            ("Events rejected", "0"),
            ("Unique source IPs", "14"),
            ("Unique users", "12"),
            ("Unique hosts", "13"),
            ("Graph nodes", "71"),
            ("Graph edges", "110"),
            ("Correlated incidents", "3"),
        ],
        widths=[3.7, 2.92], font_size=9.4)
    add_table(doc,
        ["Priority order", "Incident", "Events", "Ground-truth scenario", "Priority", "Score"],
        [
            ("1", "INC-0003", "4", "Possible exfiltration", "Critical", "100"),
            ("2", "INC-0001", "5", "Credential compromise", "Critical", "91"),
            ("3", "INC-0002", "4", "Reconnaissance", "Medium", "41"),
        ],
        widths=[0.82, 1.0, 0.65, 2.12, 1.0, 0.55], font_size=8.6)
    add_paragraph(doc,
        "Incident identifiers reflect chronological creation order, while the displayed order reflects the priority queue. "
        "The possible-exfiltration sequence ranks first because its severity total, privilege activity, sensitive-access factor, chain length, and temporal density exceed the other incident scores.")

    add_heading(doc, "12 Complexity Analysis", 1)
    add_paragraph(doc,
        "Let n be the number of events, w the maximum number of events simultaneously present in the temporal window, m the events in one incident, V the graph node count, E the graph edge count, and i the incident count.")
    add_table(doc,
        ["Operation", "Time complexity", "Additional space"],
        [
            ("CSV ingestion and indexing", "Average O(n)", "O(n)"),
            ("Hash lookup with k matches", "Average O(1 + k)", "O(k) result"),
            ("Chronological linked-list insert", "O(m)", "O(1) per node"),
            ("Sliding-window add and expire", "Amortized O(1)", "O(w)"),
            ("Candidate correlation", "O(nw)", "O(w)"),
            ("BFS or DFS", "O(V + E)", "O(V)"),
            ("Heap push or pop", "O(log i)", "O(i) total heap"),
        ],
        widths=[2.75, 1.92, 1.95], font_size=9.1)
    add_paragraph(doc,
        "The O(nw) candidate comparison is acceptable for the explainable midterm baseline. A future version can keep per-entity temporal buckets to reduce comparisons while preserving the public incident and graph interfaces.")

    add_heading(doc, "13 Current Limitations", 1)
    add_bullets(doc, [
        "The correlation policy supports three fixed event families and does not yet expose composable rule objects.",
        "Connected components provide chronological chains but do not infer causal direction between steps.",
        "The risk score is an interpretable heuristic and has not been calibrated against analyst outcomes.",
        "The dataset is controlled and small; it does not represent production alert diversity or volume.",
        "No measured throughput, latency, memory, precision, recall, or ranking-quality results exist yet.",
        "The command-line interface is the only current presentation layer.",
    ])

    add_heading(doc, "14 Planned Work", 1)
    add_numbered(doc, [
        "Reproduce the configured CMake and CTest workflow on a machine with CMake installed.",
        "Extract correlation checks into configurable rule objects and add rule-level explanations.",
        "Add directed event-chain reconstruction while retaining the chronological linked list.",
        "Create a deterministic larger dataset generator and a benchmark harness.",
        "Measure throughput, latency, memory, graph size, alert reduction, precision, recall, and ranking quality where ground truth permits.",
        "Add a dashboard only after the engine, evaluation method, and performance measurements are stable.",
    ])

    add_heading(doc, "15 Conclusion", 1)
    add_paragraph(doc,
        "The midterm implementation completes an end-to-end defensive alert-correlation pipeline in modern C++. "
        "Each required data structure has a direct role in ingestion, temporal filtering, relationship discovery, timeline storage, or prioritization. "
        "The demonstration reproduces the three planted scenarios and ranks them with visible scoring factors. "
        "The repository, tests, and progress record distinguish implemented behavior from planned work, providing a sound base for measured evaluation and more sophisticated correlation in the final project.")

    add_heading(doc, "Appendix A Repository and Execution", 1)
    add_heading(doc, "A.1 Repository Structure", 2)
    structure = (
        "include/soc    Public C++ interfaces\n"
        "src            Engine implementations and CLI\n"
        "tests          Behavioral test executable\n"
        "data           Controlled synthetic event data\n"
        "docs           Architecture, progress record, and report\n"
        "benchmarks     Reserved for measured benchmark artifacts\n"
        "scripts        Reproducible project utilities"
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    run = paragraph.add_run(structure)
    set_run_font(run, "Courier New")
    run.font.size = Pt(9)

    add_heading(doc, "A.2 Standard CMake Workflow", 2)
    commands = (
        "cmake -S . -B build\n"
        "cmake --build build\n"
        "ctest --test-dir build --output-on-failure\n"
        "./build/soc_cli data/synthetic_events.csv"
    )
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    run = paragraph.add_run(commands)
    set_run_font(run, "Courier New")
    run.font.size = Pt(9)

    add_heading(doc, "Appendix B Verified Test Summary", 1)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    output = (
        "[PASS] CSV parsing and malformed row handling\n"
        "[PASS] hash-map indexes\n"
        "[PASS] custom linked-list operations\n"
        "[PASS] sliding-window expiration\n"
        "[PASS] graph edges, BFS, and recursive DFS\n"
        "[PASS] correlation rules and incident creation\n"
        "[PASS] priority-queue ordering\n\n"
        "Tests passed: 7, failed: 0"
    )
    run = paragraph.add_run(output)
    set_run_font(run, "Courier New")
    run.font.size = Pt(8.8)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_report()
